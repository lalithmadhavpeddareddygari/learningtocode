#!/usr/bin/env python3
"""Generate Code Efficiency integration-testing proposal deck."""

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import nsmap
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

NAVY = RGBColor(0x0B, 0x2C, 0x4A)
NAVY2 = RGBColor(0x14, 0x3D, 0x63)
ORANGE = RGBColor(0xE8, 0x77, 0x22)
TEAL = RGBColor(0x0D, 0x7A, 0x86)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFFWHITE = RGBColor(0xF6, 0xF8, 0xFB)
INK = RGBColor(0x1C, 0x24, 0x30)
MUTED = RGBColor(0x5B, 0x67, 0x75)
LINE = RGBColor(0xD5, 0xDE, 0xE8)
GREEN = RGBColor(0x1F, 0x8A, 0x5B)
RED = RGBColor(0xC0, 0x39, 0x2B)
GOLD = RGBColor(0xC4, 0x8A, 0x1A)

W = Inches(13.333)
H = Inches(7.5)


def set_run(run, size=18, bold=False, color=INK, font="Calibri"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font


def add_rect(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def add_round(slide, l, t, w, h, fill):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s


def add_text(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size=size, bold=bold, color=color, font=font)
    return box


def add_bullets(slide, l, t, w, h, items, size=16, color=INK, space=10):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(space)
        p.level = 0
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size=size, color=color)
    return box


def footer(slide, n, total):
    add_rect(slide, 0, Inches(7.28), W, Inches(0.22), NAVY)
    add_text(
        slide,
        Inches(0.4),
        Inches(7.27),
        Inches(10),
        Inches(0.22),
        "Embedded Coder  ·  Code Efficiency  ·  Integration Testing Proposal  ·  Internal",
        size=10,
        color=WHITE,
    )
    add_text(
        slide,
        Inches(11.6),
        Inches(7.27),
        Inches(1.4),
        Inches(0.22),
        f"{n}  /  {total}",
        size=10,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
    )


def header_bar(slide, title, subtitle=None):
    add_rect(slide, 0, 0, W, Inches(1.15), NAVY)
    add_rect(slide, 0, Inches(1.15), W, Inches(0.08), ORANGE)
    add_text(slide, Inches(0.5), Inches(0.22), Inches(12), Inches(0.5), title, size=28, bold=True, color=WHITE)
    if subtitle:
        add_text(slide, Inches(0.5), Inches(0.68), Inches(12), Inches(0.35), subtitle, size=14, color=RGBColor(0xC9, 0xD7, 0xE5))


def card(slide, l, t, w, h, fill=WHITE):
    s = add_round(slide, l, t, w, h, fill)
    s.line.color.rgb = LINE
    s.line.width = Pt(1)
    return s


TOTAL = 20


def build():
    prs = Presentation()
    prs.slide_width = W
    prs.slide_height = H
    blank = prs.slide_layouts[6]
    n = 0

    def new():
        nonlocal n
        n += 1
        return prs.slides.add_slide(blank)

    # 1 Title
    s = new()
    add_rect(s, 0, 0, W, H, NAVY)
    add_rect(s, 0, 0, Inches(0.18), H, ORANGE)
    add_text(s, Inches(0.7), Inches(1.5), Inches(12), Inches(0.4), "QUALITY ENGINEERING  ·  MARKET RESEARCH + PROPOSAL", size=14, bold=True, color=ORANGE)
    add_text(s, Inches(0.7), Inches(2.0), Inches(12), Inches(1.6), "Integration Testing for\nEmbedded Coder Code Efficiency", size=40, bold=True, color=WHITE)
    add_text(
        s,
        Inches(0.7),
        Inches(4.0),
        Inches(11.5),
        Inches(1.2),
        "How industry teams test compiler-class products in the middle of the pyramid —\nand a practical proposal so we stop paying full codegen + SIL cost for every defect class.",
        size=18,
        color=RGBColor(0xD5, 0xE2, 0xEE),
    )
    add_text(s, Inches(0.7), Inches(6.4), Inches(11), Inches(0.4), "Audience: Code Efficiency Dev Manager  ·  Prepared by QE", size=14, color=RGBColor(0xA8, 0xBB, 0xCC))
    add_text(s, Inches(0.7), Inches(6.75), Inches(11), Inches(0.3), "Internal use  ·  September 2026", size=13, color=RGBColor(0x7F, 0x96, 0xAB))

    # 2 Why we are here
    s = new()
    header_bar(s, "The problem we are being asked to solve", "Why this research exists")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    points = [
        ("System tests are our default oracle", "Most confidence today comes from full model → codegen → compile → SIL (and sometimes profiling) workflows. That is the right test for numerical equivalence and real execution. It is the wrong test for every optimization, IR, or plumbing change."),
        ("Feedback is too late and too expensive", "Industry data is consistent: end-to-end / system tests are minutes-to-hours, high maintenance, and poor at localizing the failing component. Google, Microsoft, and compiler vendors all treat this as a design problem, not a hardware problem."),
        ("Code Efficiency is especially exposed", "Our defects are often about copies, buffer reuse, inlining, SIMD, stack, ROM/RAM, and transform interactions — not “does the model still simulate.” SIL proves functional equivalence; it does not cheaply prove that an optimization fired, stayed legal, or improved metrics."),
        ("Ask of this deck", "A market-informed definition of integration testing for our product, and a proposal we can pilot without replacing system tests."),
    ]
    for i, (title, body) in enumerate(points):
        y = Inches(1.45) + Inches(i * 1.35)
        card(s, Inches(0.45), y, Inches(12.4), Inches(1.22))
        add_rect(s, Inches(0.45), y, Inches(0.12), Inches(1.22), ORANGE if i < 3 else TEAL)
        add_text(s, Inches(0.8), y + Inches(0.12), Inches(11.8), Inches(0.35), title, size=18, bold=True, color=NAVY)
        add_text(s, Inches(0.8), y + Inches(0.45), Inches(11.8), Inches(0.7), body, size=14, color=MUTED)

    # 3 Agenda
    s = new()
    header_bar(s, "Agenda", "40–45 minutes, including discussion")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    agenda = [
        ("01", "Context", "Team charter and current test shape"),
        ("02", "Market research", "How industry defines and funds integration testing"),
        ("03", "Analogues", "Compilers (LLVM/GCC/TI) and automotive MBD (TargetLink/BTC)"),
        ("04", "Gap analysis", "What our SIL-heavy suite is actually covering — and missing"),
        ("05", "Proposal", "A Code Efficiency integration layer, CI policy, and pilot"),
        ("06", "Decision", "What I need from you to start"),
    ]
    for i, (num, title, desc) in enumerate(agenda):
        col = i % 3
        row = i // 3
        x = Inches(0.5) + Inches(col * 4.2)
        y = Inches(1.7) + Inches(row * 2.4)
        card(s, x, y, Inches(3.95), Inches(2.1))
        add_text(s, x + Inches(0.25), y + Inches(0.25), Inches(3.4), Inches(0.5), num, size=28, bold=True, color=ORANGE)
        add_text(s, x + Inches(0.25), y + Inches(0.85), Inches(3.4), Inches(0.4), title, size=20, bold=True, color=NAVY)
        add_text(s, x + Inches(0.25), y + Inches(1.3), Inches(3.4), Inches(0.6), desc, size=14, color=MUTED)

    # 4 Team
    s = new()
    header_bar(s, "Team context: Code Efficiency", "What we own, and therefore what we must test")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    cols = [
        (ORANGE, "Generated-code efficiency", "Optimizations that reduce ROM, RAM, copies, stack, and execution time: buffer reuse, expression folding, inlining, loop transforms, SIMD, parallel for, saturation/protection stripping."),
        (TEAL, "Codegen workflows & infra", "Configuration objectives, Code Generation Advisor, static metrics, reports, traceability, and productivity tooling that sit on the codegen path."),
        (NAVY, "Engineering quality", "Processes that keep codegen products shippable: regression signal, localization, and preventing “optimization vs correctness” regressions."),
    ]
    for i, (color, title, body) in enumerate(cols):
        x = Inches(0.45) + Inches(i * 4.25)
        card(s, x, Inches(1.55), Inches(4.05), Inches(4.0))
        add_rect(s, x, Inches(1.55), Inches(4.05), Inches(0.12), color)
        add_text(s, x + Inches(0.25), Inches(1.9), Inches(3.55), Inches(1.0), title, size=20, bold=True, color=NAVY)
        add_text(s, x + Inches(0.25), Inches(3.0), Inches(3.55), Inches(2.2), body, size=15, color=MUTED)
    add_text(
        s,
        Inches(0.5),
        Inches(5.8),
        Inches(12.3),
        Inches(1.1),
        "Implication for QE: a test that only asks “SIL matches Normal mode” is necessary but not sufficient. Efficiency work also needs oracles for structure (did the transform fire?), metrics (did ROM/RAM/stack move the right way?), and legality (did we keep numerical / language semantics?).",
        size=15,
        color=INK,
    )

    # 5 Current shape
    s = new()
    header_bar(s, "Current test shape: an inverted pyramid", "Not a people problem — a layering problem")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    layers = [
        (Inches(3.6), Inches(1.55), Inches(6.1), Inches(1.15), RED, WHITE, "SYSTEM  ·  Full codegen + compile + SIL / PIL  ·  minutes–hours  ·  high localization cost"),
        (Inches(4.3), Inches(2.85), Inches(4.7), Inches(1.15), GOLD, WHITE, "INTEGRATION  ·  Thin / informal today  ·  the missing middle"),
        (Inches(5.0), Inches(4.15), Inches(3.3), Inches(1.15), GREEN, WHITE, "UNIT  ·  Isolated transforms  ·  fast"),
    ]
    for l, t, w, h, fill, tc, txt in layers:
        add_round(s, l, t, w, h, fill)
        add_text(s, l, t + Inches(0.35), w, Inches(0.5), txt, size=13, bold=True, color=tc, align=PP_ALIGN.CENTER)
    card(s, Inches(0.4), Inches(5.5), Inches(12.5), Inches(1.5))
    add_bullets(
        s,
        Inches(0.6),
        Inches(5.62),
        Inches(12.1),
        Inches(1.3),
        [
            "Industry target mix (Google SWE book): ~80% small/unit, ~15% medium/integration, ~5% large/E2E. Typical industry pyramid: 60–70 / 20–30 / 5–10.",
            "A SIL-heavy suite behaves like an ice-cream cone: slow, brittle, and expensive to diagnose. That is the #1 anti-pattern called out by Google, Microsoft, Fowler, and ISTQB-aligned guides.",
            "We should keep SIL. We should stop using it as the first test we write for every Code Efficiency change.",
        ],
        size=14,
        space=6,
    )

    # 6 Industry definition
    s = new()
    header_bar(s, "What industry means by integration testing", "A distinct defect class — not a smaller system test")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    headers = ["", "Unit", "Integration", "System / E2E"]
    rows = [
        ["Question", "Does this function do the right thing in isolation?", "Do two or more real components honor their interface?", "Does the whole product journey still work?"],
        ["Defect class", "Logic, algorithm, local invariant", "Contracts, data flow, miswired stages, config", "Emergent behavior, environment, UX of the workflow"],
        ["Typical time", "ms–seconds", "seconds–low minutes", "minutes–hours"],
        ["Owner mix", "Dev-primary", "Dev + QE together", "QE-primary, sparse"],
        ["Oracle", "Assert on values / mocks", "Assert on boundary artifacts", "Assert on customer-visible outcome"],
    ]
    # table-like cards
    col_w = [Inches(2.0), Inches(3.35), Inches(3.55), Inches(3.55)]
    x0 = Inches(0.4)
    y0 = Inches(1.5)
    for j, htxt in enumerate(headers):
        x = x0 + sum(col_w[:j], Inches(0))
        add_rect(s, x, y0, col_w[j], Inches(0.45), NAVY)
        add_text(s, x + Inches(0.08), y0 + Inches(0.08), col_w[j] - Inches(0.1), Inches(0.32), htxt, size=13, bold=True, color=WHITE)
    for i, row in enumerate(rows):
        y = y0 + Inches(0.45) + Inches(i * 0.85)
        bg = WHITE if i % 2 == 0 else RGBColor(0xEE, 0xF3, 0xF8)
        for j, cell in enumerate(row):
            x = x0 + sum(col_w[:j], Inches(0))
            add_rect(s, x, y, col_w[j], Inches(0.85), bg)
            add_text(
                s,
                x + Inches(0.08),
                y + Inches(0.12),
                col_w[j] - Inches(0.12),
                Inches(0.65),
                cell,
                size=12,
                bold=(j == 0),
                color=NAVY if j == 0 else INK,
            )
    add_text(
        s,
        Inches(0.45),
        Inches(6.65),
        Inches(12.4),
        Inches(0.4),
        "Fowler / Microsoft: keep integration tests narrow. If a test needs the whole stack, it is a system test wearing an integration-test name.",
        size=13,
        color=MUTED,
    )

    # 7 Consensus
    s = new()
    header_bar(s, "Industry consensus (2024–2026)", "Same advice, different companies")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    quotes = [
        ("Google (SWE book)", "Classify tests by size (small / medium / large) and by scope. Aim ~80 / 15 / 5. Write the smallest test that can fail for that behavior. Small = one process; medium = one machine; large = anywhere."),
        ("Microsoft (Well-Architected)", "Unit on every commit. Integration when components interact. E2E only for the few critical journeys. Gate pipelines so a change cannot proceed until the cheaper layer is green."),
        ("ISTQB / test-level guides", "Integration owns interfaces. System owns the fully integrated product against requirements. Different owners, different environments, different oracles."),
        ("Compiler industry", "Do not wait for whole-program execution to find a bad transform. Lit/FileCheck-style tests prove a stage produced the expected IR/code. Whole programs are a second, slower net."),
    ]
    for i, (who, txt) in enumerate(quotes):
        y = Inches(1.5) + Inches(i * 1.3)
        card(s, Inches(0.45), y, Inches(12.4), Inches(1.18))
        add_text(s, Inches(0.7), y + Inches(0.12), Inches(12), Inches(0.3), who, size=14, bold=True, color=ORANGE)
        add_text(s, Inches(0.7), y + Inches(0.45), Inches(12), Inches(0.65), txt, size=14, color=INK)

    # 8 LLVM
    s = new()
    header_bar(s, "Analogue 1: how compiler vendors test", "LLVM, GCC, TI — the closest technical cousins")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    items = [
        ("LLVM unit tests", "C++ unit tests in llvm/unittests. Must pass before commit. Fast, hermetic, algorithm-level."),
        ("LLVM regression / lit + FileCheck", "Tiny IR or source snippets. RUN the relevant tool (opt, llc, clang). CHECK that generated IR/asm contains the expected pattern. This is the industry gold standard for codegen integration tests. Seconds, not minutes. Localizes to a pass."),
        ("LLVM test-suite (whole programs)", "Real programs with reference outputs, compile-time, code size, runtime. Closer to our SIL + metrics suite. Not the first line of defense."),
        ("TI / commercial compilers", "Conformance suites (Plum Hall, SuperTest) + in-house kernels + application benchmarks + bug regressions. Kernels check performance and correctness of hot loops — a model for our optimization micro-fixtures."),
        ("Fuzz / differential (GraphicsFuzz, Csmith, CompFuzzCI)", "Random programs + compare compilers or optimization levels. Amazon/Dafny work even puts short fuzz campaigns on PRs. Complementary later; not the first pilot."),
    ]
    for i, (t, b) in enumerate(items):
        y = Inches(1.42) + Inches(i * 1.05)
        add_round(s, Inches(0.45), y, Inches(0.55), Inches(0.85), TEAL)
        add_text(s, Inches(0.45), y + Inches(0.25), Inches(0.55), Inches(0.4), str(i + 1), size=18, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        card(s, Inches(1.15), y, Inches(11.7), Inches(0.95))
        add_text(s, Inches(1.35), y + Inches(0.08), Inches(11.3), Inches(0.3), t, size=15, bold=True, color=NAVY)
        add_text(s, Inches(1.35), y + Inches(0.4), Inches(11.3), Inches(0.5), b, size=13, color=MUTED)

    # 9 Automotive
    s = new()
    header_bar(s, "Analogue 2: automotive MBD & codegen QA", "How the market around Embedded Coder tests generated code")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    cards = [
        ("Customers (our users)", "MIL → SIL → PIL back-to-back, often ISO 26262. Tools: Simulink Test, SIL/PIL Manager, BTC EmbeddedTester (TargetLink and Embedded Coder), EverTest. Their SIL is product verification of generated code vs model. Ours is product verification of the code generator. Same mechanism, different system under test."),
        ("dSPACE TargetLink + BTC", "Harness extraction from model/code architecture, requirements tests, automated test-vector generation, MIL/SIL/PIL, Jenkins. Migration suites compare tool versions — a pattern we should copy internally when we change optimizers."),
        ("Ford / CRL case (public BTC)", "Hundreds of reusable library routines, multiple MATLAB releases, Jenkins. Tests sit at the library-function boundary, not a full vehicle SIL. That is integration-scale for generated/replaced code."),
        ("MathWorks product guidance", "Advisor + static metrics for efficiency; SIL/PIL profiling for timing. The product itself already separates static structure from dynamic execution. Our test strategy should mirror that split."),
    ]
    for i, (t, b) in enumerate(cards):
        col = i % 2
        row = i // 2
        x = Inches(0.4) + Inches(col * 6.45)
        y = Inches(1.5) + Inches(row * 2.55)
        card(s, x, y, Inches(6.25), Inches(2.35))
        add_rect(s, x, y, Inches(0.12), Inches(2.35), ORANGE)
        add_text(s, x + Inches(0.35), y + Inches(0.2), Inches(5.7), Inches(0.45), t, size=16, bold=True, color=NAVY)
        add_text(s, x + Inches(0.35), y + Inches(0.7), Inches(5.7), Inches(1.5), b, size=13, color=MUTED)

    # 10 Mapping
    s = new()
    header_bar(s, "Mapping industry layers onto Embedded Coder", "Name the layers in our vocabulary so the team can use them")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    map_rows = [
        ("Google size", "Our name", "What it actually runs", "When"),
        ("Small", "Unit", "One transform, cost model, or API with fakes. No Simulink engine if avoidable.", "Every CL / presubmit"),
        ("Medium", "Integration", "Minimal model or IR fixture → real codegen pipeline through N stages. Assert IR, C, metrics, or compile. Stop before SIL unless the defect class needs execution.", "Every CL if < few minutes; else gated + merge"),
        ("Large", "System", "Representative models, full codegen, compile, SIL/PIL, equivalence, profiling.", "Nightly / pre-release / selected smoke"),
        ("Huge / soak", "Qualification & fuzz", "Conformance, random models, version-to-version migration, long-tail targets.", "Scheduled / milestone"),
    ]
    colw = [Inches(2.0), Inches(2.2), Inches(5.7), Inches(2.5)]
    x0 = Inches(0.45)
    y0 = Inches(1.5)
    for j, htxt in enumerate(map_rows[0]):
        x = x0 + sum(colw[:j], Inches(0))
        add_rect(s, x, y0, colw[j], Inches(0.5), NAVY)
        add_text(s, x + Inches(0.1), y0 + Inches(0.1), colw[j] - Inches(0.15), Inches(0.32), htxt, size=13, bold=True, color=WHITE)
    for i, row in enumerate(map_rows[1:]):
        y = y0 + Inches(0.5) + Inches(i * 1.15)
        bg = WHITE if i % 2 == 0 else RGBColor(0xEE, 0xF3, 0xF8)
        for j, cell in enumerate(row):
            x = x0 + sum(colw[:j], Inches(0))
            add_rect(s, x, y, colw[j], Inches(1.15), bg)
            add_text(s, x + Inches(0.1), y + Inches(0.15), colw[j] - Inches(0.18), Inches(0.9), cell, size=13, bold=(j < 2), color=INK)

    # 11 Gap
    s = new()
    header_bar(s, "Gap analysis for Code Efficiency", "SIL is a great system test. It is an expensive integration test.")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    left = [
        "A copy was not eliminated, but numbers still match.",
        "An optimization did not fire (wrong IR, wrong setting).",
        "Two passes interact and undo each other.",
        "Static metrics / report plumbing broke.",
        "Generated C still compiles but structure regressed.",
        "A config objective mapped to the wrong Details flags.",
    ]
    right = [
        "Numerical mismatch vs Normal / MIL.",
        "Runtime crash, timeout, or stack overflow in SIL.",
        "Profiling numbers on a representative model.",
        "Target / toolchain / PIL issues.",
        "Emergent behavior across referenced models.",
        "Customer-visible codegen + verify workflow.",
    ]
    card(s, Inches(0.4), Inches(1.5), Inches(6.2), Inches(4.55))
    add_rect(s, Inches(0.4), Inches(1.5), Inches(6.2), Inches(0.55), RGBColor(0xC0, 0x39, 0x2B))
    add_text(s, Inches(0.6), Inches(1.58), Inches(5.8), Inches(0.4), "SIL is a weak / late detector for…", size=16, bold=True, color=WHITE)
    add_bullets(s, Inches(0.65), Inches(2.2), Inches(5.7), Inches(3.6), left, size=15, space=8)
    card(s, Inches(6.8), Inches(1.5), Inches(6.1), Inches(4.55))
    add_rect(s, Inches(6.8), Inches(1.5), Inches(6.1), Inches(0.55), GREEN)
    add_text(s, Inches(7.0), Inches(1.58), Inches(5.7), Inches(0.4), "Keep SIL as the detector for…", size=16, bold=True, color=WHITE)
    add_bullets(s, Inches(7.05), Inches(2.2), Inches(5.6), Inches(3.6), right, size=15, space=8)
    add_text(
        s,
        Inches(0.45),
        Inches(6.2),
        Inches(12.4),
        Inches(0.8),
        "Proposal principle: match the oracle to the defect class. Efficiency bugs are mostly structural and metric. Correctness bugs still need execution. Do not force one oracle to serve both.",
        size=15,
        color=INK,
    )

    # 12 Proposal architecture
    s = new()
    header_bar(s, "Proposal: a four-layer strategy for this team", "Integration is Layer 2 — new investment, not a rename of SIL")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    layers2 = [
        ("L1  Unit", "Existing + grow", "Transform unit tests, cost models, config mapping tables, report parsers. No model load when possible.", GREEN),
        ("L2  Integration  ★", "New focus", "Four families: (A) FileCheck-style codegen snapshots, (B) pipeline-stage contracts, (C) static metric contracts, (D) host-compile kernels. No SIL by default.", ORANGE),
        ("L3  System smoke", "Shrink & curate", "A small, stable SIL pack per optimization area. Must be diagnosable. Runs on merge / nightly, not every local save.", TEAL),
        ("L4  System deep", "Keep, schedule", "Broad models, PIL, profiling, migration vs last release. Nightly / weekly. Failures file bugs; they do not gate every CL.", NAVY),
    ]
    for i, (t, tag, b, c) in enumerate(layers2):
        y = Inches(1.45) + Inches(i * 1.3)
        add_round(s, Inches(0.45), y, Inches(2.6), Inches(1.15), c)
        add_text(s, Inches(0.55), y + Inches(0.25), Inches(2.4), Inches(0.4), t, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        add_text(s, Inches(0.55), y + Inches(0.65), Inches(2.4), Inches(0.35), tag, size=12, color=WHITE, align=PP_ALIGN.CENTER)
        card(s, Inches(3.25), y, Inches(9.6), Inches(1.15))
        add_text(s, Inches(3.5), y + Inches(0.25), Inches(9.2), Inches(0.7), b, size=15, color=INK)

    # 13 Four families
    s = new()
    header_bar(s, "Layer 2 in detail: four integration families", "Start with A+C. Add B and D once the harness exists.")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    fam = [
        ("A. Snapshot / FileCheck", "Tiny models or IR. Generate C (or dump IR). Match stable patterns: function signature, reuse of a buffer, absence of a memcpy, presence of a SIMD intrinsic, inlined callee gone. Auto-update golden with review, like LLVM update_*_test_checks."),
        ("B. Pipeline contracts", "Assert stage handoff: model compile → IR → optimized IR → emit → compile. Example: after buffer-reuse pass, live ranges share storage; emitter does not reintroduce a temp. Failure names the pass, not “SIL mismatch.”"),
        ("C. Metric contracts", "Use our own static code metrics (the product already computes ROM/RAM/stack/lines/globals). Assert bounds or deltas vs a baseline fixture. Catches efficiency regressions without executing a step."),
        ("D. Kernel compile+run", "For a handful of hot loops: generate, compile on host, run a tiny driver with known I/O. Cheaper than SIL Manager + SDI compare; still catches semantic breaks in the optimizer. Closest to LLVM kernels / TI in-house kernels."),
    ]
    for i, (t, b) in enumerate(fam):
        col = i % 2
        row = i // 2
        x = Inches(0.4) + Inches(col * 6.45)
        y = Inches(1.5) + Inches(row * 2.55)
        card(s, x, y, Inches(6.25), Inches(2.35))
        add_rect(s, x, y, Inches(6.25), Inches(0.5), NAVY if i else ORANGE)
        add_text(s, x + Inches(0.25), y + Inches(0.1), Inches(5.8), Inches(0.35), t, size=16, bold=True, color=WHITE)
        add_text(s, x + Inches(0.25), y + Inches(0.7), Inches(5.8), Inches(1.5), b, size=13, color=INK)

    # 14 Example
    s = new()
    header_bar(s, "Worked example: buffer reuse change", "Same feature, three test designs — only one belongs in SIL")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    ex = [
        ("Unit (minutes to write, ms to run)", "Given two live ranges that do not overlap, allocator assigns the same storage. Table-driven cases for sizes, dimensions, Reusable storage class, Model-block outputs."),
        ("Integration A+C (seconds)", "3-block fixture model. Generate. CHECK: one local for the path, no extra memcpy. Metric: RAM globals ≤ baseline, stack of step() ≤ bound. Compile the C. Stop."),
        ("System SIL (keep 1–2)", "One representative referenced-model model, Normal vs SIL equivalence, plus static metrics trend. Proves we did not break numerics or the customer workflow. Not 40 permutations of the same reuse rule."),
    ]
    for i, (t, b) in enumerate(ex):
        y = Inches(1.5) + Inches(i * 1.55)
        card(s, Inches(0.45), y, Inches(12.4), Inches(1.4))
        add_text(s, Inches(0.7), y + Inches(0.15), Inches(12), Inches(0.35), t, size=18, bold=True, color=NAVY)
        add_text(s, Inches(0.7), y + Inches(0.55), Inches(12), Inches(0.7), b, size=15, color=MUTED)
    add_text(s, Inches(0.5), Inches(6.25), Inches(12.3), Inches(0.7), "Rule of thumb for QE reviews: if the expected failure message cannot name the pass or the metric, the test is too big.", size=15, color=INK)

    # 15 CI
    s = new()
    header_bar(s, "When tests run (industry CI policy, applied here)", "Speed of the suite is a product requirement, not a luxury")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    ci = [
        ("Local / presubmit", "L1 + affected L2 (A/C) for the optimization area. Target: minutes, hermetic, no SIL."),
        ("Merge / gated CI", "Full L2 for the component + a tiny L3 smoke (few SIL). Fail the merge."),
        ("Nightly", "L3 remaining + L4 models, profiling on a fixed set, last-release migration on a fixture pack."),
        ("Milestone / release", "Broader targets, PIL sample, qualification evidence if needed."),
    ]
    for i, (t, b) in enumerate(ci):
        x = Inches(0.4) + Inches(i * 3.2)
        card(s, x, Inches(1.55), Inches(3.05), Inches(3.3))
        add_rect(s, x, Inches(1.55), Inches(3.05), Inches(0.12), ORANGE)
        add_text(s, x + Inches(0.15), Inches(1.85), Inches(2.75), Inches(0.8), t, size=16, bold=True, color=NAVY)
        add_text(s, x + Inches(0.15), Inches(2.7), Inches(2.75), Inches(1.9), b, size=13, color=MUTED)
    card(s, Inches(0.4), Inches(5.05), Inches(12.5), Inches(1.9))
    add_text(s, Inches(0.65), Inches(5.2), Inches(12), Inches(0.35), "Selection, not “run everything”", size=16, bold=True, color=NAVY)
    add_text(
        s,
        Inches(0.65),
        Inches(5.6),
        Inches(12),
        Inches(1.15),
        "Google TAP lesson: test volume × commit rate will not scale if every test is large. Use area ownership (tests live next to the pass), and impacted-test selection where the build system allows. A change in buffer reuse should not wait on SIMD SIL. Flaky L3/L4 tests are disabled aggressively — a red nightly with no owner is not a quality process.",
        size=14,
        color=INK,
    )

    # 16 Pilot
    s = new()
    header_bar(s, "Pilot plan (QE-led, Dev-paired)", "Prove the middle layer on one optimization family before a platform bet")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    phases = [
        ("Phase 0  Align", "Pick one family (recommend: buffer reuse / copies, or SIMD). Inventory today’s SIL tests that exist only to detect “optimization didn’t fire.” Agree oracles: patterns + static metrics."),
        ("Phase 1  Harness", "Minimal runner: load fixture → slbuild/codegen → dump C/IR/metrics → match. One command, CI-able, artifact on fail. No SIL in the harness."),
        ("Phase 2  Seed suite", "20–40 fixtures covering the family, including known historical bugs. Convert 5–10 existing slow system tests into L2 + keep 2 SIL sentinels."),
        ("Phase 3  Policy", "Presubmit = L2 for that folder. Measure: median time, defects found before SIL, false-golden updates. Then replicate the harness to the next family."),
    ]
    for i, (t, b) in enumerate(phases):
        y = Inches(1.45) + Inches(i * 1.25)
        add_round(s, Inches(0.45), y, Inches(2.5), Inches(1.1), NAVY if i else ORANGE)
        add_text(s, Inches(0.55), y + Inches(0.35), Inches(2.3), Inches(0.45), t, size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        card(s, Inches(3.15), y, Inches(9.7), Inches(1.1))
        add_text(s, Inches(3.4), y + Inches(0.2), Inches(9.3), Inches(0.75), b, size=14, color=INK)

    # 17 Metrics
    s = new()
    header_bar(s, "How we will know it worked", "Pilot success metrics — discuss and freeze in Phase 0")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    mets = [
        ("Feedback time", "L2 pack for the family runs in minutes on a standard engineer machine / CI shard. Contrast with current SIL pack time."),
        ("Defects shifted left", "Count of bugs (or injected faults) caught by L2 that previously needed SIL. Target: majority of “didn’t fire / extra copy / metric regression” class."),
        ("Localization", "Median time from red to blamed pass/file. L2 failures should name fixture + check, not a 20-minute SIL log."),
        ("SIL diet", "Number of redundant system tests retired or demoted, without increasing escaped customer defects in that family."),
        ("Golden hygiene", "Rate of snapshot updates that are justified vs noisy. If goldens churn on unrelated edits, tighten CHECK patterns (LLVM lesson)."),
        ("Dev adoption", "New optimization CLs in the family include L2 tests by default in review. QE reviews test level, not only coverage of SIL."),
    ]
    for i, (t, b) in enumerate(mets):
        col = i % 3
        row = i // 3
        x = Inches(0.4) + Inches(col * 4.25)
        y = Inches(1.5) + Inches(row * 2.55)
        card(s, x, y, Inches(4.05), Inches(2.35))
        add_text(s, x + Inches(0.2), y + Inches(0.2), Inches(3.65), Inches(0.6), t, size=16, bold=True, color=ORANGE)
        add_text(s, x + Inches(0.2), y + Inches(0.85), Inches(3.65), Inches(1.3), b, size=13, color=INK)

    # 18 Risks
    s = new()
    header_bar(s, "Risks and how we contain them", "Integration tests fail in predictable ways — plan for them")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    risks = [
        ("Goldens become change detectors", "CHECK patterns, not whole-file diffs. Review golden updates like code. Auto-update scripts with diffs in the CL."),
        ("We accidentally rebuild SIL in disguise", "Harness forbids SIL/PIL APIs in L2. Time budget per test. QE checklist: “can this fail without executing generated code?”"),
        ("False confidence, skip SIL entirely", "Keep sentinel SIL per family. Never let L2 be the only numerical oracle for a pass that can change values."),
        ("Fixture models bit-rot", "Own fixtures next to the pass. Run L2 on every CL that touches that directory. Treat fixtures as product code."),
        ("Tooling cost higher than expected", "Phase 1 is a thin MATLAB/Python runner, not a new product. Reuse static metrics and existing codegen APIs."),
        ("Org inertia (“we already have system tests”)", "Pilot on one family with before/after numbers. Do not ask for a suite rewrite up front."),
    ]
    for i, (t, b) in enumerate(risks):
        col = i % 2
        row = i // 2
        x = Inches(0.4) + Inches(col * 6.45)
        y = Inches(1.45) + Inches(row * 1.75)
        card(s, x, y, Inches(6.25), Inches(1.6))
        add_text(s, x + Inches(0.25), y + Inches(0.15), Inches(5.8), Inches(0.4), t, size=15, bold=True, color=NAVY)
        add_text(s, x + Inches(0.25), y + Inches(0.6), Inches(5.8), Inches(0.85), b, size=13, color=MUTED)

    # 19 Ask
    s = new()
    header_bar(s, "Recommendation and ask", "What I need from you to start")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    add_text(s, Inches(0.5), Inches(1.5), Inches(12.3), Inches(0.7), "Recommend we fund a Code Efficiency integration-test layer, piloted on one optimization family, rather than scaling SIL further.", size=18, bold=True, color=NAVY)
    asks = [
        "Endorse the four-layer vocabulary (Unit / Integration / System smoke / System deep) so reviews have a shared language.",
        "Name a Dev counterpart for Phase 0–1 (harness + first fixtures) and pick the family (buffer reuse/copies vs SIMD).",
        "Allow QE to propose retiring or demoting SIL cases once L2 covers the same defect class — with sentinel SIL retained.",
        "Protect time for the harness as engineering work, not “extra test writing on the side of feature QE.”",
        "If the pilot’s metrics look good, replicate the pattern; if not, we stop — this is a bounded experiment, not a religion.",
    ]
    add_bullets(s, Inches(0.55), Inches(2.3), Inches(12.2), Inches(3.5), asks, size=16, space=12)
    card(s, Inches(0.45), Inches(5.85), Inches(12.4), Inches(1.1))
    add_text(s, Inches(0.7), Inches(6.1), Inches(12), Inches(0.7), "Decision needed today: approve Phase 0–1 on one family. Everything else waits on those numbers.", size=16, bold=True, color=TEAL)

    # 20 Sources
    s = new()
    header_bar(s, "Sources (selected)", "So the research is auditable")
    footer(s, n, TOTAL)
    add_rect(s, 0, Inches(1.23), W, Inches(6.05), OFFWHITE)
    src = [
        "Google, Software Engineering at Google, Ch. 11 Testing Overview & Ch. 14 Larger Testing — small/medium/large, ~80/15/5 pyramid.",
        "Microsoft Learn, Azure Well-Architected — testing strategies; pyramid; E2E kept small; pipeline gates.",
        "Martin Fowler on integration-test scope; ISTQB test levels (component / integration / system / acceptance).",
        "LLVM Testing Infrastructure Guide; FileCheck; llvm-test-suite (whole programs vs lit regressions).",
        "TI compiler validation (Plum Hall, SuperTest, kernels, applications, regressions).",
        "Donaldson et al., GraphicsFuzz / gfauto; CompFuzzCI (Dafny/Amazon) — fuzz as a complement to regression.",
        "MathWorks Embedded Coder: SIL/PIL Manager, static code metrics, code efficiency techniques, execution profiling.",
        "BTC EmbeddedTester / dSPACE TargetLink: MIL/SIL/PIL, migration suites, Ford CRL+Jenkins case; EverTest for Simulink B2B.",
    ]
    add_bullets(s, Inches(0.55), Inches(1.5), Inches(12.2), Inches(5.5), src, size=14, space=8)

    out = "/workspace/integration-testing-proposal/Integration_Testing_Market_Research.pptx"
    prs.save(out)
    print("Wrote", out, "slides", n)


if __name__ == "__main__":
    build()
